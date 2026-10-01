"""
Conferência de Serviços Prestados: DecWeb x Portal Nacional (Emitidas)
=====================================================================

Antes de enviar a declaração à prefeitura (ação oficial e irreversível), compara o que o DecWeb vai
declarar com o que o Portal Nacional mostra como emitido pelo cliente na mesma competência:

  - lado DecWeb: relatório "Notas Serviços Prestados" (Nota Nacional) que a Fase 1 baixa em zip e extrai na
    pasta 002 ARQUIVOS MUNICIPAIS do mês, na convenção do escritório:
        {n}_{apelido}_PortoAlegre_MM.AAAA_NFSE_ServPrestados.zip
        {n}_{apelido}_PortoAlegre_MM.AAAA_NFSE_ServPrestados\\<inscrição>_AAAAMM_NFSE_ServPrestados.xlsx
    O .xlsx tem o cabeçalho "ISSQNdec - Relação de Serviços Prestados"; a tabela começa na linha cuja 1ª coluna
    é "Número NFSE" e termina na linha "TOTAL".
  - lado Nacional: planilha Emitidas exportada do Portal Nacional, em 003 ARQUIVOS PORTAL NACIONAL do mês, com
    o nome da razão social ("<RAZÃO SOCIAL> - Emitidas.xlsx"). É localizada pelo CNPJ do PRESTADOR dentro da
    planilha (coluna "CNPJ/CPF Prestador"), nunca pelo nome do arquivo. Filtrada pela coluna Competência, sem
    notas canceladas.

Situações devolvidas por conferir_cliente():
  ok            DecWeb e Nacional batem (ou os dois sem movimento)
  divergencia   notas ou valores diferentes (tolerância de R$ 0,01)
  sem_emitidas  não há planilha Emitidas do cliente para conferir

O número da nota no DecWeb vem com prefixo ("202610000000001342") e no Nacional é o número puro (1342): casa
pelos 6 últimos dígitos do DecWeb.

Pode rodar sozinho, sem site:
    python conferencia_prestados.py 09/2026 47
"""

import csv
import io
import re
import sys
import zipfile
from pathlib import Path

import openpyxl

import configuracao

TOLERANCIA = 0.01
MESES_PASTA_MUN = configuracao.MESES_PASTA  # compatibilidade com o pacote original


def _para_float(v) -> float:
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    return float(str(v).strip().replace("R$", "").replace(" ", "").replace(".", "").replace(",", ".")
                 if "," in str(v) else str(v).strip())


def _so_digitos(v) -> str:
    return re.sub(r"\D", "", str(v or ""))


# ----------------------------------------------------------------------
# LEITURA DAS PLANILHAS
# ----------------------------------------------------------------------

def ler_decweb_prestados(origem) -> dict:
    """origem: caminho do .xlsx ou bytes. Devolve {'competencia', 'notas': {numero:int -> {valor, situacao}}}
    das notas NÃO canceladas."""
    wb = openpyxl.load_workbook(io.BytesIO(origem) if isinstance(origem, bytes) else origem, data_only=True)
    ws = wb.worksheets[0]
    notas: dict[int, dict] = {}
    competencia = ""
    cabecalho = None
    for row in ws.iter_rows(values_only=True):
        primeiro = str(row[0]).strip() if row and row[0] is not None else ""
        if primeiro.lower().startswith("competência"):
            competencia = str(row[1] or "").strip()
        if primeiro == "Número NFSE":
            cabecalho = [str(c or "").strip() for c in row]
            continue
        if cabecalho is None:
            continue
        if not primeiro or any(str(c).strip().upper() == "TOTAL" for c in row if c is not None):
            break  # fim da tabela (linha de total)
        ix_valor = cabecalho.index("Valor Serviços")
        ix_sit = next((i for i, c in enumerate(cabecalho) if c.startswith("Situação")), None)
        situacao = str(row[ix_sit] or "") if ix_sit is not None else ""
        if "cancel" in situacao.lower():
            continue
        digitos = _so_digitos(primeiro)
        numero = int(digitos[-6:])
        notas[numero] = {"valor": _para_float(row[ix_valor]), "situacao": situacao}
    return {"competencia": competencia, "notas": notas}


def ler_nacional_emitidas(caminho: Path, competencia: str) -> dict:
    """Notas da competência (MM/AAAA) da planilha Emitidas, sem canceladas."""
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        linhas = ws.iter_rows(values_only=True)
        cab = [str(c or "").strip() for c in next(linhas)]
        i_num = cab.index("Número NFS-e")
        i_comp = cab.index("Competência")
        i_val = cab.index("Valor do Serviço (R$)")
        i_sit = [i for i, c in enumerate(cab) if c in ("Situação NFS-e", "Situação")]
        notas: dict[int, dict] = {}
        for r in linhas:
            if not r or r[i_num] in (None, ""):
                continue
            if not str(r[i_comp] or "").startswith(competencia):
                continue
            situacao = " ".join(str(r[i] or "") for i in i_sit)
            if "cancel" in situacao.lower():
                continue
            notas[int(_so_digitos(r[i_num]))] = {"valor": _para_float(r[i_val]), "situacao": situacao}
        return {"notas": notas}
    finally:
        wb.close()


def cnpj_prestador_do_arquivo(caminho: Path) -> str:
    """CNPJ (só dígitos) do prestador, lido da 1ª linha de dados da planilha Emitidas. '' se não achar."""
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        linhas = ws.iter_rows(values_only=True)
        cab = [str(c or "").strip() for c in next(linhas, [])]
        if "CNPJ/CPF Prestador" not in cab:
            return ""
        i = cab.index("CNPJ/CPF Prestador")
        for r in linhas:
            if r and r[i] not in (None, ""):
                return _so_digitos(r[i])
        return ""
    finally:
        wb.close()


def indexar_emitidas_por_cnpj(pasta_nacional: Path) -> dict[str, Path]:
    """{cnpj -> planilha Emitidas} da pasta 003 do mês. Se houver mais de uma por CNPJ (ex.: 'Emitidas' e
    'Emitidas - Completa'), prefere a que NÃO é 'Completa' (a mesma planilha base usada nas demais etapas)."""
    indice: dict[str, Path] = {}
    if not pasta_nacional.exists():
        return indice
    for arq in sorted(pasta_nacional.glob("*Emitidas*.xlsx")):
        if arq.name.startswith("~$"):
            continue
        try:
            cnpj = cnpj_prestador_do_arquivo(arq)
        except Exception:
            continue
        if not cnpj:
            continue
        atual = indice.get(cnpj)
        if atual is None or ("completa" in atual.name.lower() and "completa" not in arq.name.lower()):
            indice[cnpj] = arq
    return indice


# ----------------------------------------------------------------------
# COMPARAÇÃO
# ----------------------------------------------------------------------

def comparar(decweb: dict, nacional: dict) -> dict:
    d, n = decweb["notas"], nacional["notas"]
    so_decweb = sorted(set(d) - set(n))
    so_nacional = sorted(set(n) - set(d))
    valor_diferente = [
        (num, d[num]["valor"], n[num]["valor"])
        for num in sorted(set(d) & set(n))
        if abs(d[num]["valor"] - n[num]["valor"]) > TOLERANCIA
    ]
    total_d = round(sum(x["valor"] for x in d.values()), 2)
    total_n = round(sum(x["valor"] for x in n.values()), 2)
    return {
        "ok": not (so_decweb or so_nacional or valor_diferente),
        "qtd_decweb": len(d), "qtd_nacional": len(n),
        "total_decweb": total_d, "total_nacional": total_n,
        "so_decweb": [(x, d[x]["valor"]) for x in so_decweb],
        "so_nacional": [(x, n[x]["valor"]) for x in so_nacional],
        "valor_diferente": valor_diferente,
    }


def _fmt(v: float) -> str:
    return f"R$ {v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def resumir(res: dict) -> str:
    """Texto curto pro log/Painel."""
    if res["ok"]:
        return f"conferido: {res['qtd_decweb']} nota(s), total {_fmt(res['total_decweb'])}"
    partes = [f"DecWeb {res['qtd_decweb']} nota(s) {_fmt(res['total_decweb'])} x Nacional "
              f"{res['qtd_nacional']} nota(s) {_fmt(res['total_nacional'])}"]
    if res["so_nacional"]:
        partes.append("só no Nacional: " + ", ".join(f"NFS-e {n} ({_fmt(v)})" for n, v in res["so_nacional"][:8]))
    if res["so_decweb"]:
        partes.append("só no DecWeb: " + ", ".join(f"NFS-e {n} ({_fmt(v)})" for n, v in res["so_decweb"][:8]))
    if res["valor_diferente"]:
        partes.append("valor diferente: " + ", ".join(
            f"NFS-e {n} (DecWeb {_fmt(a)} x Nacional {_fmt(b)})" for n, a, b in res["valor_diferente"][:8]))
    return "DIVERGÊNCIA — " + "; ".join(partes)


# ----------------------------------------------------------------------
# LOCALIZAÇÃO DOS ARQUIVOS DE UM CLIENTE
# ----------------------------------------------------------------------

def nome_base_prestados(cliente: dict, mes: str, ano: str) -> str:
    return f"{cliente['numero']}_{cliente['apelido']}_PortoAlegre_{mes}.{ano}_NFSE_ServPrestados"


def achar_decweb_prestados(pasta_mun: Path, cliente: dict, mes: str, ano: str):
    """Lê o Prestados do DecWeb: primeiro a subpasta extraída, depois o zip. None se nenhum existir."""
    base = nome_base_prestados(cliente, mes, ano)
    pasta = pasta_mun / base
    if pasta.is_dir():
        xlsx = sorted(p for p in pasta.glob("*.xlsx") if not p.name.startswith("~$"))
        if xlsx:
            return ler_decweb_prestados(str(xlsx[0]))
    arq_zip = pasta_mun / f"{base}.zip"
    if arq_zip.exists():
        with zipfile.ZipFile(arq_zip) as z:
            xlsx = [n for n in z.namelist() if n.lower().endswith(".xlsx")]
            if xlsx:
                return ler_decweb_prestados(z.read(xlsx[0]))
    return None


def achar_nacional_emitidas(pasta_nac: Path, cliente: dict, mes: str, ano: str, competencia: str,
                            indice: dict | None = None):
    """Planilha Emitidas do cliente, achada pelo CNPJ dentro do arquivo (ou, sem CNPJ no cadastro, pelo padrão
    antigo '{n}_*_MM.AAAA_Emitidas.xlsx'). None se não existir."""
    arq = None
    cnpj = _so_digitos(cliente.get("cnpj"))
    if cnpj:
        arq = (indice if indice is not None else indexar_emitidas_por_cnpj(pasta_nac)).get(cnpj)
    if arq is None and pasta_nac.exists():
        antigos = sorted(pasta_nac.glob(f"{cliente['numero']}_*_{mes}.{ano}_Emitidas.xlsx"))
        arq = antigos[0] if antigos else None
    return ler_nacional_emitidas(arq, competencia) if arq else None


def conferir_cliente(cliente: dict, competencia: str, cfg: dict | None = None) -> tuple[str, str]:
    """Confere um cliente (dict com 'numero', 'apelido' e, de preferência, 'cnpj') na competência MM/AAAA.
    Devolve (situacao, texto): 'ok' | 'divergencia' | 'sem_emitidas'."""
    cfg = cfg or configuracao.carregar()
    mes, ano = competencia.split("/")
    decweb = achar_decweb_prestados(configuracao.pasta_municipais(cfg, mes), cliente, mes, ano)
    nacional = achar_nacional_emitidas(configuracao.pasta_nacional(cfg, mes), cliente, mes, ano, competencia)

    if nacional is None:
        if decweb is not None and decweb["notas"]:
            return "sem_emitidas", (f"não há planilha Emitidas do Portal Nacional para este cliente — "
                                    f"não foi possível conferir as {len(decweb['notas'])} nota(s) do DecWeb")
        return "sem_emitidas", "não há planilha Emitidas do Portal Nacional (e o DecWeb não tem notas de Prestados)"

    if decweb is None:
        if not nacional["notas"]:
            return "ok", "sem movimento nos dois lados (sem Prestados no DecWeb e Emitidas sem notas da competência)"
        return "divergencia", (f"o Nacional tem {len(nacional['notas'])} nota(s) da competência, "
                               "mas não há Prestados do DecWeb para este cliente (rode a Fase 1)")

    res = comparar(decweb, nacional)
    return ("ok" if res["ok"] else "divergencia"), resumir(res)


if __name__ == "__main__":
    comp = sys.argv[1]
    outros = sys.argv[2:]
    arq = Path(__file__).parent / "config" / "clientes.csv"
    with open(arq, encoding="utf-8-sig") as f:
        todos = [{k.strip(): (v or "").strip() for k, v in r.items()} for r in csv.DictReader(f)]
    for c in [c for c in todos if not outros or c["numero"] in outros]:
        sit, txt = conferir_cliente(c, comp)
        print(f"{c['numero']:>3} {c['apelido']:<28} {sit:<12} {txt}")
