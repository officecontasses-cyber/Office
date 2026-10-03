"""
Conciliação DecWeb x Portal Nacional (Prestados x Emitidas, Tomados x Recebidas, declaração e guia)
==================================================================================================

Roda sozinho, SEM abrir site, a partir dos arquivos que os dois robôs já gravaram no Drive:

  002 ARQUIVOS MUNICIPAIS  (robô DecWeb)
      {n}_{apelido}_PortoAlegre_MM.AAAA_NFSE_ServPrestados[.zip | \\<inscrição>_AAAAMM_NFSE_ServPrestados.xlsx]
      {n}_{apelido}_PortoAlegre_MM.AAAA_NFSE_ServTomados  [idem]
      {n}_{apelido}_PortoAlegre_MM.AAAA_DeclaraçãoMensal.pdf   (opcional, precisa do pacote pypdf)
      {n}_{apelido}_PortoAlegre_MM.AAAA_ISSQN.pdf              (a guia; opcional)
  003 ARQUIVOS PORTAL NACIONAL  (robô Portal Nacional)
      {n}_{apelido}_PortoAlegre_MM.AAAA_Emitidas.xlsx
      {n}_{apelido}_PortoAlegre_MM.AAAA_Recebidas.xlsx

O que é conferido, por cliente:
  PRESTADOS   notas do DecWeb x Emitidas do Portal (competência MM/AAAA, sem canceladas): quantidade, valor de cada
              nota, ISS e ISS retido. As notas são casadas pela CHAVE DE ACESSO (igual nos dois lados); se um lado
              não tiver chave, pelo número da nota (6 últimos dígitos no DecWeb).
  TOMADOS     idem, DecWeb Tomados x Recebidas do Portal.
  DECLARAÇÃO  receita bruta e "total a recolher" do recibo da prefeitura contra o Prestados do DecWeb e contra a
              guia (valor e vencimento). Só se o pypdf estiver instalado.

Situações por lado: ok | atencao | divergencia | sem_portal | sem_dados
  divergencia  nota só de um lado, valor da nota diferente ou ISS retido diferente (tolerância R$ 0,01)
  atencao      bate no essencial, mas há algo para olhar (ISS diferente, nota com valor zero, notas de outra
               competência no Portal, declaração diferente do Prestados...)
  sem_portal   o DecWeb tem arquivo e o Portal não (não baixado, ou o robô não grava planilha com 0 notas)
  sem_dados    nenhum dos dois lados tem arquivo

Uso (na pasta do robô DecWeb, usa config/configuracao.ini e config/clientes.csv):
    python conciliacao.py 09/2026                    # todos os clientes do clientes.csv que tenham arquivos
    python conciliacao.py 09/2026 155 16 238         # só alguns (pelo número)
    python conciliacao.py 09/2026 --csv conciliacao_09_2026.csv
    python conciliacao.py 09/2026 --detalhe          # lista nota a nota o que divergiu
    python conciliacao.py x --ler-pdf "caminho\\155_..._DeclaraçãoMensal.pdf"   # confere a leitura de um PDF

NUNCA lê nem imprime usuário/senha do clientes.csv: só 'numero' e 'apelido'.
"""

import argparse
import csv
import io
import re
import sys
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

import openpyxl

import configuracao

TOLERANCIA = 0.01
TIPOS_DECWEB = {"prestados": "ServPrestados", "tomados": "ServTomados"}
TIPOS_PORTAL = {"prestados": "Emitidas", "tomados": "Recebidas"}


# ----------------------------------------------------------------------
# UTILITÁRIOS
# ----------------------------------------------------------------------

def so_digitos(v) -> str:
    return re.sub(r"\D", "", str(v or ""))


def para_float(v) -> float:
    """Número vindo da planilha: float/int, ou texto em formato BR ("1.234,56") ou US ("1,234.56")."""
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("R$", "").replace(" ", "")
    if not s or s in ("-", "--"):
        return 0.0
    if "," in s and "." in s:
        # o último separador é o decimal
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "," in s:
        if re.fullmatch(r"-?\d{1,3}(,\d{3})+", s):
            s = s.replace(",", "")          # 1,234 -> milhar
        else:
            s = s.replace(",", ".")        # 12,5 / 1234,56 -> decimal
    elif "." in s and re.fullmatch(r"-?\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")              # 1.234 / 1.234.567 -> milhar (dinheiro não tem 3 decimais)
    try:
        return float(s)
    except ValueError:
        return 0.0


def brl(v: float) -> str:
    return f"R$ {v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def normaliza_competencia(v) -> str:
    """'09/2026', '09/2026 Original' ou data -> 'MM/AAAA'."""
    if isinstance(v, (datetime, date)):
        return f"{v.month:02d}/{v.year}"
    m = re.search(r"(\d{2})/(\d{4})", str(v or ""))
    return f"{m.group(1)}/{m.group(2)}" if m else ""


def _idx(cabecalho: list[str], *nomes: str):
    """Índice da 1ª coluna cujo nome (sem espaços nas pontas, sem ponto final, minúsculo) bate com algum dos nomes."""
    norm = [c.strip().rstrip(".").casefold() for c in cabecalho]
    for nome in nomes:
        n = nome.strip().rstrip(".").casefold()
        if n in norm:
            return norm.index(n)
    return None


# ----------------------------------------------------------------------
# MODELO
# ----------------------------------------------------------------------

@dataclass
class Nota:
    numero: int
    valor: float
    iss: float = 0.0
    iss_retido: float = 0.0
    situacao: str = ""
    chave: str = ""
    competencia: str = ""

    @property
    def cancelada(self) -> bool:
        s = self.situacao.lower()
        return "cancel" in s or "substitu" in s


@dataclass
class Lado:
    """Notas de um lado (DecWeb ou Portal) já filtradas: competência certa e sem canceladas."""
    notas: list[Nota] = field(default_factory=list)
    canceladas: int = 0
    outras_competencias: list[Nota] = field(default_factory=list)  # só Portal: notas de outro mês no arquivo

    @property
    def qtd(self) -> int:
        return len(self.notas)

    @property
    def total(self) -> float:
        return round(sum(n.valor for n in self.notas), 2)

    @property
    def iss(self) -> float:
        return round(sum(n.iss for n in self.notas), 2)

    @property
    def retido(self) -> float:
        return round(sum(n.iss_retido for n in self.notas), 2)


# ----------------------------------------------------------------------
# LEITURA DAS PLANILHAS
# ----------------------------------------------------------------------

def _abre(origem):
    return openpyxl.load_workbook(io.BytesIO(origem) if isinstance(origem, (bytes, bytearray)) else origem,
                                  read_only=True, data_only=True)


def ler_decweb(origem) -> Lado:
    """Planilha do DecWeb (ISSQNdec - Relação de Serviços Prestados/Tomados). origem: caminho ou bytes."""
    wb = _abre(origem)
    try:
        ws = wb.worksheets[0]
        cab = None
        lado = Lado()
        for row in ws.iter_rows(values_only=True):
            row = list(row)
            primeiro = str(row[0]).strip() if row and row[0] is not None else ""
            if primeiro == "Número NFSE":
                cab = [str(c or "").strip() for c in row]
                i_val = _idx(cab, "Valor Serviços")
                i_iss = _idx(cab, "Val. ISS")
                i_ret = _idx(cab, "Val. ISS ret.")
                i_sit = _idx(cab, "Situação.", "Situação")
                i_chave = _idx(cab, "Chave Acesso")
                continue
            if cab is None:
                continue
            if not primeiro or any(str(c).strip().upper() == "TOTAL" for c in row if c is not None):
                break  # linha de total
            situacao = str(row[i_sit] or "") if i_sit is not None else ""
            nota = Nota(
                numero=int(so_digitos(primeiro)[-6:] or 0),
                valor=para_float(row[i_val]) if i_val is not None else 0.0,
                iss=para_float(row[i_iss]) if i_iss is not None else 0.0,
                iss_retido=para_float(row[i_ret]) if i_ret is not None else 0.0,
                situacao=situacao,
                chave=so_digitos(row[i_chave]) if i_chave is not None else "",
            )
            if nota.cancelada:
                lado.canceladas += 1
            else:
                lado.notas.append(nota)
        return lado
    finally:
        wb.close()


def ler_portal(caminho, competencia: str) -> Lado:
    """Planilha Emitidas/Recebidas do Portal Nacional. Fica só com a competência MM/AAAA e sem canceladas;
    notas de outras competências vão para `outras_competencias` (só informativo)."""
    wb = _abre(caminho)
    try:
        ws = wb.worksheets[0]
        linhas = ws.iter_rows(values_only=True)
        cab = [str(c or "").strip() for c in next(linhas)]
        i_num = _idx(cab, "Número NFS-e", "Relação Número NFS-e")
        i_comp = _idx(cab, "Competência")
        i_val = _idx(cab, "Valor do Serviço (R$)")
        i_iss = _idx(cab, "Valor do ISSQN (R$)")
        i_ret = _idx(cab, "Retenção ISSQN")
        i_sit = [i for i, c in enumerate(cab) if c.strip() in ("Situação NFS-e", "Situação")]
        i_chave = _idx(cab, "Chave NFS-e")
        faltando = [n for n, i in (("Número NFS-e", i_num), ("Competência", i_comp), ("Valor do Serviço (R$)", i_val)) if i is None]
        if faltando:
            raise ValueError(f"{caminho}: faltam as colunas {', '.join(faltando)} (não parece uma planilha do Portal Nacional)")
        lado = Lado()
        for r in linhas:
            if not r or r[i_num] in (None, "") or not so_digitos(r[i_num]):
                continue  # linhas de total/resumo
            iss = para_float(r[i_iss]) if i_iss is not None else 0.0
            ret_txt = str(r[i_ret] or "") if i_ret is not None else ""
            retido = iss if ret_txt.strip()[:1] in ("2", "3") else 0.0   # "2 - Retido pelo Tomador"
            nota = Nota(
                numero=int(so_digitos(r[i_num])),
                valor=para_float(r[i_val]),
                iss=iss,
                iss_retido=retido,
                situacao=" ".join(str(r[i] or "") for i in i_sit),
                chave=so_digitos(r[i_chave]) if i_chave is not None else "",
                competencia=normaliza_competencia(r[i_comp]),
            )
            if nota.competencia != competencia:
                lado.outras_competencias.append(nota)
            elif nota.cancelada:
                lado.canceladas += 1
            else:
                lado.notas.append(nota)
        return lado
    finally:
        wb.close()


# ----------------------------------------------------------------------
# COMPARAÇÃO
# ----------------------------------------------------------------------

@dataclass
class Comparacao:
    decweb: Lado
    portal: Lado
    so_decweb: list[Nota] = field(default_factory=list)
    so_portal: list[Nota] = field(default_factory=list)
    valor_diferente: list[tuple[Nota, Nota]] = field(default_factory=list)
    retido_diferente: list[tuple[Nota, Nota]] = field(default_factory=list)
    iss_diferente: list[tuple[Nota, Nota]] = field(default_factory=list)
    zeradas: list[Nota] = field(default_factory=list)

    @property
    def situacao(self) -> str:
        if self.so_decweb or self.so_portal or self.valor_diferente or self.retido_diferente:
            return "divergencia"
        if self.iss_diferente or self.zeradas or self.portal.outras_competencias or self.portal.canceladas != self.decweb.canceladas:
            return "atencao"
        return "ok"


def comparar(decweb: Lado, portal: Lado) -> Comparacao:
    comp = Comparacao(decweb=decweb, portal=portal)
    pend_d = list(decweb.notas)
    pend_p = list(portal.notas)
    pares: list[tuple[Nota, Nota]] = []

    # 1ª passada: chave de acesso
    por_chave = {n.chave: n for n in pend_p if n.chave}
    resto_d = []
    usados = set()
    for d in pend_d:
        p = por_chave.get(d.chave) if d.chave else None
        if p is not None and id(p) not in usados:
            pares.append((d, p))
            usados.add(id(p))
        else:
            resto_d.append(d)
    resto_p = [p for p in pend_p if id(p) not in usados]

    # 2ª passada: número da nota (6 últimos dígitos) para o que sobrou
    por_num: dict[int, Nota] = {}
    for p in resto_p:
        por_num.setdefault(p.numero % 1_000_000, p)
    sobra_d = []
    usados2 = set()
    for d in resto_d:
        p = por_num.get(d.numero % 1_000_000)
        if p is not None and id(p) not in usados2:
            pares.append((d, p))
            usados2.add(id(p))
        else:
            sobra_d.append(d)
    sobra_p = [p for p in resto_p if id(p) not in usados2]

    comp.so_decweb = sorted(sobra_d, key=lambda n: n.numero)
    comp.so_portal = sorted(sobra_p, key=lambda n: n.numero)
    for d, p in pares:
        if abs(d.valor - p.valor) > TOLERANCIA:
            comp.valor_diferente.append((d, p))
        if abs(d.iss_retido - p.iss_retido) > TOLERANCIA:
            comp.retido_diferente.append((d, p))
        if abs(d.iss - p.iss) > TOLERANCIA:
            comp.iss_diferente.append((d, p))
    comp.zeradas = [n for n in decweb.notas if abs(n.valor) <= TOLERANCIA]
    return comp


# ----------------------------------------------------------------------
# DECLARAÇÃO E GUIA (PDF)
# ----------------------------------------------------------------------

_VALOR = r"(\d{1,3}(?:\.\d{3})*,\d{2})"


def extrair_declaracao(texto: str) -> dict:
    """Receita bruta (soma das escriturações), total a recolher e data/hora do recibo, do texto da DeclaraçãoMensal."""
    t = re.sub(r"\s+", " ", texto or "")
    receitas = [para_float(x) for x in re.findall(r"Base de C[áa]lculo Receita Bruta " + _VALOR, t)]
    totais = re.findall(r"Total Geral A Recolher " + _VALOR, t)
    recibo = re.search(r"DECLARA[ÇC][ÃA]O RECEBIDA EM (\d{2}/\d{2}/\d{4}) [àa]s (\d{2}:\d{2}:\d{2})", t)
    return {
        "receita": round(sum(receitas), 2) if receitas else None,
        "a_recolher": para_float(totais[-1]) if totais else None,
        "recibo": f"{recibo.group(1)} {recibo.group(2)}" if recibo else None,
    }


def extrair_guia(texto: str) -> dict:
    """Valor a pagar e vencimento impressos na guia de ISSQN."""
    t = re.sub(r"\s+", " ", texto or "")
    valor = re.search(r"VALOR A PAGAR R\$ " + _VALOR, t)
    venc = re.search(r"N[ÃA]O RECEBER ESTA GUIA AP[ÓO]S (\d{2}/\d{2}/\d{4})", t)
    return {"valor": para_float(valor.group(1)) if valor else None, "vencimento": venc.group(1) if venc else None}


def texto_pdf(caminho: Path) -> str | None:
    """Texto do PDF, ou None se o pypdf não estiver instalado."""
    try:
        from pypdf import PdfReader
    except ImportError:
        return None
    return "\n".join((p.extract_text() or "") for p in PdfReader(str(caminho)).pages)


# ----------------------------------------------------------------------
# LOCALIZAÇÃO DOS ARQUIVOS DE UM CLIENTE
# ----------------------------------------------------------------------

def nome_base(cliente: dict, mes: str, ano: str) -> str:
    return f"{cliente['numero']}_{cliente['apelido']}_PortoAlegre_{mes}.{ano}"


def achar_decweb(pasta_mun: Path, cliente: dict, mes: str, ano: str, tipo: str):
    """Lado DecWeb (pasta extraída primeiro, depois o zip). None se o cliente não tem o arquivo."""
    base = f"{nome_base(cliente, mes, ano)}_NFSE_{TIPOS_DECWEB[tipo]}"
    pasta = pasta_mun / base
    if pasta.is_dir():
        xlsx = sorted(p for p in pasta.glob("*.xlsx") if not p.name.startswith("~$"))
        if xlsx:
            return ler_decweb(str(xlsx[0]))
    arq_zip = pasta_mun / f"{base}.zip"
    if arq_zip.exists():
        with zipfile.ZipFile(arq_zip) as z:
            xlsx = [n for n in z.namelist() if n.lower().endswith(".xlsx")]
            if xlsx:
                return ler_decweb(z.read(xlsx[0]))
    return None


def achar_portal(pasta_nac: Path, cliente: dict, mes: str, ano: str, tipo: str, competencia: str):
    """Lado Portal Nacional. None se não há planilha (não baixada, ou o robô não grava planilha com 0 notas)."""
    kind = TIPOS_PORTAL[tipo]
    arq = pasta_nac / f"{nome_base(cliente, mes, ano)}_{kind}.xlsx"
    if not arq.exists():
        achados = sorted(p for p in pasta_nac.glob(f"{cliente['numero']}_*_{mes}.{ano}_{kind}.xlsx")
                         if not p.name.startswith("~$")) if pasta_nac.exists() else []
        arq = achados[0] if achados else None
    return ler_portal(arq, competencia) if arq else None


# ----------------------------------------------------------------------
# CONCILIAÇÃO DE UM CLIENTE
# ----------------------------------------------------------------------

def conciliar_lado(decweb: Lado | None, portal: Lado | None) -> tuple[str, Comparacao | None, str]:
    """(situacao, comparacao, texto curto) de um lado (prestados ou tomados)."""
    if decweb is None and portal is None:
        return "sem_dados", None, "sem arquivo nos dois lados"
    if portal is None:
        return ("sem_portal", None,
                f"DecWeb {decweb.qtd} nota(s) {brl(decweb.total)}; sem planilha do Portal "
                "(não baixada, ou o Portal não teve notas)")
    if decweb is None:
        if not portal.notas:
            return "ok", None, "sem notas nos dois lados"
        return ("divergencia", None,
                f"Portal {portal.qtd} nota(s) {brl(portal.total)} mas sem arquivo do DecWeb (rode a Fase 1)")
    comp = comparar(decweb, portal)
    base = (f"DecWeb {decweb.qtd} {brl(decweb.total)} (ISS {brl(decweb.iss)}, retido {brl(decweb.retido)}) x "
            f"Portal {portal.qtd} {brl(portal.total)} (ISS {brl(portal.iss)}, retido {brl(portal.retido)})")
    extras = []
    if comp.so_decweb:
        extras.append(f"só no DecWeb: {len(comp.so_decweb)}")
    if comp.so_portal:
        extras.append(f"só no Portal: {len(comp.so_portal)}")
    if comp.valor_diferente:
        extras.append(f"valor diferente: {len(comp.valor_diferente)}")
    if comp.retido_diferente:
        extras.append(f"ISS retido diferente: {len(comp.retido_diferente)}")
    if comp.iss_diferente:
        extras.append(f"ISS diferente (informativo): {len(comp.iss_diferente)}")
    if comp.zeradas:
        extras.append(f"nota(s) com valor zero: {len(comp.zeradas)}")
    if portal.canceladas or decweb.canceladas:
        extras.append(f"canceladas Portal {portal.canceladas} / DecWeb {decweb.canceladas}")
    if portal.outras_competencias:
        comps = sorted({n.competencia for n in portal.outras_competencias})
        extras.append(f"Portal tem {len(portal.outras_competencias)} nota(s) de outra competência ({', '.join(comps)}), ignoradas")
    return comp.situacao, comp, base + ("; " + "; ".join(extras) if extras else "")


def conciliar_declaracao(cliente: dict, mes: str, ano: str, pasta_mun: Path, prest_decweb: Lado | None) -> tuple[str, str]:
    """Confere o recibo da DeclaraçãoMensal e a guia. ('ok'|'atencao'|'sem_dados'|'divergencia', texto)."""
    base = nome_base(cliente, mes, ano)
    pdf_decl = pasta_mun / f"{base}_DeclaraçãoMensal.pdf"
    pdf_guia = pasta_mun / f"{base}_ISSQN.pdf"
    if not pdf_decl.exists():
        return "sem_dados", "sem DeclaraçãoMensal.pdf (declaração ainda não enviada?)"
    txt = texto_pdf(pdf_decl)
    if txt is None:
        return "sem_dados", "pypdf não instalado (pip install pypdf): PDFs não conferidos"
    decl = extrair_declaracao(txt)
    partes, nivel = [], "ok"
    if decl["receita"] is None or decl["a_recolher"] is None:
        # texto do PDF fora do formato esperado: nunca ficar em silêncio
        partes.append("NÃO consegui ler " + " e ".join(
            n for n, v in (("a receita bruta", decl["receita"]), ("o total a recolher", decl["a_recolher"])) if v is None)
            + " do PDF (confira com --ler-pdf)")
        nivel = "atencao"
    if decl["recibo"]:
        partes.append(f"recibo {decl['recibo']}")
    else:
        partes.append("recibo NÃO encontrado no PDF")
        nivel = "atencao"
    if decl["receita"] is not None and prest_decweb is not None:
        if abs(decl["receita"] - prest_decweb.total) > TOLERANCIA:
            partes.append(f"receita da declaração {brl(decl['receita'])} x Prestados do DecWeb {brl(prest_decweb.total)}")
            nivel = "atencao"
        else:
            partes.append(f"receita {brl(decl['receita'])} = Prestados do DecWeb")
    elif decl["receita"] is not None:
        partes.append(f"receita {brl(decl['receita'])} (sem Prestados do DecWeb para comparar)")
    a_recolher = decl["a_recolher"]
    if a_recolher is not None:
        if pdf_guia.exists():
            guia = extrair_guia(texto_pdf(pdf_guia) or "")
            if guia["valor"] is None:
                partes.append(f"a recolher {brl(a_recolher)}; valor da guia não lido")
                nivel = "atencao"
            elif abs(guia["valor"] - a_recolher) > TOLERANCIA:
                partes.append(f"a recolher {brl(a_recolher)} x guia {brl(guia['valor'])}")
                nivel = "divergencia"
            else:
                partes.append(f"a recolher {brl(a_recolher)} = guia, venc. {guia['vencimento'] or '?'}")
        elif a_recolher > TOLERANCIA:
            partes.append(f"a recolher {brl(a_recolher)} mas SEM guia ISSQN.pdf")
            nivel = "atencao"
        else:
            partes.append("a recolher R$ 0,00 (sem guia)")
    return nivel, "; ".join(partes)


def conciliar_cliente(cliente: dict, competencia: str, cfg: dict | None = None) -> dict:
    cfg = cfg or configuracao.carregar()
    mes, ano = competencia.split("/")
    mun = configuracao.pasta_municipais(cfg, mes)
    nac = configuracao.pasta_nacional(cfg, mes)
    out = {"cliente": cliente, "lados": {}}
    prest_decweb = None
    for tipo in ("prestados", "tomados"):
        dec = achar_decweb(mun, cliente, mes, ano, tipo)
        por = achar_portal(nac, cliente, mes, ano, tipo, competencia)
        if tipo == "prestados":
            prest_decweb = dec
        sit, comp, texto = conciliar_lado(dec, por)
        out["lados"][tipo] = {"situacao": sit, "texto": texto, "comparacao": comp, "decweb": dec, "portal": por}
    sit, texto = conciliar_declaracao(cliente, mes, ano, mun, prest_decweb)
    out["lados"]["declaracao"] = {"situacao": sit, "texto": texto, "comparacao": None, "decweb": None, "portal": None}
    return out


# ----------------------------------------------------------------------
# SAÍDA
# ----------------------------------------------------------------------

ROTULO = {"prestados": "PRESTADOS ", "tomados": "TOMADOS   ", "declaracao": "DECLARAÇÃO"}
ORDEM = {"divergencia": 3, "atencao": 2, "sem_portal": 1, "ok": 0, "sem_dados": 0}


def detalhe(comp: Comparacao) -> list[str]:
    linhas = []
    for n in comp.so_decweb:
        linhas.append(f"      só no DecWeb : NFS-e {n.numero} {brl(n.valor)} ISS {brl(n.iss)} retido {brl(n.iss_retido)}")
    for n in comp.so_portal:
        linhas.append(f"      só no Portal : NFS-e {n.numero} {brl(n.valor)} ISS {brl(n.iss)} retido {brl(n.iss_retido)}")
    for d, p in comp.valor_diferente:
        linhas.append(f"      valor        : NFS-e {p.numero} DecWeb {brl(d.valor)} x Portal {brl(p.valor)}")
    for d, p in comp.retido_diferente:
        linhas.append(f"      ISS retido   : NFS-e {p.numero} DecWeb {brl(d.iss_retido)} x Portal {brl(p.iss_retido)}")
    for d, p in comp.iss_diferente:
        linhas.append(f"      ISS          : NFS-e {p.numero} DecWeb {brl(d.iss)} x Portal {brl(p.iss)}")
    for n in comp.zeradas:
        linhas.append(f"      valor zero   : NFS-e {n.numero} ({n.situacao or 'sem situação'})")
    return linhas


def imprimir(resultados: list[dict], com_detalhe: bool) -> None:
    for r in resultados:
        c = r["cliente"]
        print(f"\n{c['numero']:>4} {c['apelido']}")
        for tipo in ("prestados", "tomados", "declaracao"):
            lado = r["lados"][tipo]
            print(f"   {ROTULO[tipo]} {lado['situacao'].upper():<11} {lado['texto']}")
            if com_detalhe and lado["comparacao"] is not None:
                for ln in detalhe(lado["comparacao"]):
                    print(ln)
    contagem: dict[str, int] = {}
    for r in resultados:
        pior = max((l["situacao"] for l in r["lados"].values()), key=lambda s: ORDEM[s])
        contagem[pior] = contagem.get(pior, 0) + 1
    print("\nRESUMO: " + ", ".join(f"{k} {v}" for k, v in sorted(contagem.items(), key=lambda kv: -ORDEM[kv[0]])))


def gravar_csv(resultados: list[dict], caminho: str) -> None:
    with open(caminho, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["numero", "apelido", "lado", "situacao", "qtd_decweb", "total_decweb", "iss_decweb", "retido_decweb",
                    "qtd_portal", "total_portal", "iss_portal", "retido_portal", "detalhe"])
        for r in resultados:
            c = r["cliente"]
            for tipo in ("prestados", "tomados", "declaracao"):
                lado = r["lados"][tipo]
                d, p = lado["decweb"], lado["portal"]
                w.writerow([
                    c["numero"], c["apelido"], tipo, lado["situacao"],
                    d.qtd if d else "", d.total if d else "", d.iss if d else "", d.retido if d else "",
                    p.qtd if p else "", p.total if p else "", p.iss if p else "", p.retido if p else "",
                    lado["texto"],
                ])


def carregar_clientes(numeros: list[str]) -> list[dict]:
    arq = Path(__file__).parent / "config" / "clientes.csv"
    if not arq.exists():
        sys.exit(f"Não encontrei {arq}.")
    with open(arq, encoding="utf-8-sig", newline="") as f:
        todos = []
        for linha in csv.DictReader(f):
            c = {(k or "").strip(): (v or "").strip() for k, v in linha.items() if k}
            if c.get("numero"):
                todos.append({"numero": c["numero"], "apelido": c.get("apelido", "")})  # nunca carrega usuário/senha
    return [c for c in todos if not numeros or c["numero"] in numeros]


def main(argv=None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    ap = argparse.ArgumentParser(description="Conciliação DecWeb x Portal Nacional.")
    ap.add_argument("competencia", help="MM/AAAA (com --ler-pdf, qualquer valor)")
    ap.add_argument("clientes", nargs="*", help="números dos clientes (se omitido, todos do clientes.csv)")
    ap.add_argument("--csv", help="grava o resultado neste arquivo CSV (separador ;)")
    ap.add_argument("--detalhe", action="store_true", help="lista nota a nota o que divergiu")
    ap.add_argument("--ler-pdf", metavar="ARQUIVO", help="só mostra o texto e os campos lidos de uma DeclaraçãoMensal.pdf ou "
                    "guia ISSQN.pdf (para conferir a leitura); nesse modo a competência é ignorada")
    args = ap.parse_args(argv)
    if args.ler_pdf:
        if not Path(args.ler_pdf).is_file():
            print(f"Não encontrei o arquivo {args.ler_pdf}")
            return 1
        txt = texto_pdf(Path(args.ler_pdf))
        if txt is None:
            print("pypdf não instalado: pip install pypdf")
            return 1
        print(txt)
        print("\n--- campos lidos ---")
        print("declaração:", extrair_declaracao(txt))
        print("guia      :", extrair_guia(txt))
        return 0
    if not re.fullmatch(r"(0[1-9]|1[0-2])/\d{4}", args.competencia):
        ap.error("competência deve ser MM/AAAA")
    cfg = configuracao.carregar()
    resultados = []
    for c in carregar_clientes(args.clientes):
        r = conciliar_cliente(c, args.competencia, cfg)
        if args.clientes or any(l["situacao"] != "sem_dados" for l in r["lados"].values()):
            resultados.append(r)
    if not resultados:
        print("Nenhum cliente com arquivos nessa competência.")
        return 1
    imprimir(resultados, args.detalhe)
    if args.csv:
        gravar_csv(resultados, args.csv)
        print(f"CSV gravado em {args.csv}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
