"""Testes da conciliação DecWeb x Portal Nacional (offline, planilhas sintéticas no layout real)."""
import io
import shutil
import zipfile

import openpyxl
import pytest

import conciliacao as co
import configuracao

CLIENTE = {"numero": "238", "apelido": "ReatHolding"}
CHAVE_A = "43149022234914125000149000000000049326094479430148"
CHAVE_B = "43149022234914125000149000000000049426090775266030"
CHAVE_C = "43214361239430957000103000000000099426098376608832"
CHAVE_D = "43149022246766362000199000000000012126090814083657"


# ---------------------------------------------------------------- planilhas sintéticas

def xlsx_decweb(notas, tipo="Tomados", competencia="09/2026 Original"):
    """notas: [(numero_com_prefixo, valor, iss, iss_retido, situacao, chave)] no layout do ISSQNdec."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([f"ISSQNdec - Relação de Serviços {tipo}"])
    ws.append(["Inscrição Municipal:", "714567-2-4"])
    ws.append(["Competência:", competencia])
    ws.append(["Nome da Empresa:", "EMPRESA TESTE"])
    ws.append(["CNPJ:", "49.252.432/0001-89"])
    ws.append(["Número NFSE", "Data Em. Nota", "Valor Serviços", "Val. Ded.", "Base Cálculo", "Aliq. (%)", "Val. ISS",
               "Val. ISS ret.", "Situação.", "Chave Acesso"])
    for num, valor, iss, ret, sit, chave in notas:
        ws.append([num, "01/09/2026", valor, 0, valor, 5, iss, ret, sit, chave])
    ws.append(["", "", "TOTAL", "TOTAL"])
    ws.append(["", "", sum(n[1] for n in notas), 0])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def xlsx_portal(notas):
    """notas: [(numero, competencia, valor, iss, retencao, situacao, chave)] no layout do Portal Nacional."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Relação Número NFS-e", "Data Geração", "Competência", "Valor do Serviço (R$)", "Situação NFS-e",
               "Valor do ISSQN (R$)", "Retenção ISSQN", "Situação", "Chave NFS-e"])
    for num, comp, valor, iss, ret, sit, chave in notas:
        ws.append([num, "01/09/2026", comp, valor, "100 - NFS-e Gerada", iss, ret, sit, chave])
    ws.append([None, None, None, "TOTAL (n notas)"])          # linhas de total/resumo do Portal
    ws.append(["Resumo Situação", "Qtd", "Valor (R$)"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture()
def pastas():
    cfg = configuracao.carregar()
    mun = configuracao.pasta_municipais(cfg, "09")
    nac = configuracao.pasta_nacional(cfg, "09")
    mun.mkdir(parents=True, exist_ok=True)
    nac.mkdir(parents=True, exist_ok=True)
    for pasta in (mun, nac):
        for f in pasta.iterdir():
            shutil.rmtree(f) if f.is_dir() else f.unlink()
    return mun, nac


def grava_decweb(mun, tipo, conteudo, como_zip=False):
    base = f"{co.nome_base(CLIENTE, '09', '2026')}_NFSE_{co.TIPOS_DECWEB[tipo]}"
    if como_zip:
        with zipfile.ZipFile(mun / f"{base}.zip", "w") as z:
            z.writestr("714567_202609.xlsx", conteudo)
    else:
        (mun / base).mkdir(exist_ok=True)
        (mun / base / "714567_202609.xlsx").write_bytes(conteudo)


def grava_portal(nac, tipo, conteudo):
    (nac / f"{co.nome_base(CLIENTE, '09', '2026')}_{co.TIPOS_PORTAL[tipo]}.xlsx").write_bytes(conteudo)


# ---------------------------------------------------------------- utilitários

@pytest.mark.parametrize("entrada,esperado", [
    (1234.5, 1234.5), (None, 0.0), ("", 0.0), ("2,527.86", 2527.86), ("2.527,86", 2527.86), ("1234,56", 1234.56),
    ("12,5", 12.5), ("1,234", 1234.0), ("1.234", 1234.0), ("1.234.567", 1234567.0), ("275.00", 275.0),
    ("R$ 842,02", 842.02), ("-", 0.0),
])
def test_para_float_aceita_formato_br_e_us(entrada, esperado):
    assert co.para_float(entrada) == pytest.approx(esperado)


def test_normaliza_competencia():
    assert co.normaliza_competencia("09/2026 Original") == "09/2026"
    import datetime
    assert co.normaliza_competencia(datetime.date(2026, 9, 1)) == "09/2026"
    assert co.normaliza_competencia(None) == ""


# ---------------------------------------------------------------- leitores

def test_leitor_decweb_ignora_canceladas_e_para_no_total():
    lado = co.ler_decweb(xlsx_decweb([
        ("202610000000000493", 9249.04, 462.45, 462.45, "Normal", CHAVE_A),
        ("202610000000000494", 100.0, 5.0, 0, "Cancelada", CHAVE_B),
    ]))
    assert lado.qtd == 1 and lado.canceladas == 1
    n = lado.notas[0]
    assert (n.numero, n.valor, n.iss, n.iss_retido, n.chave) == (493, 9249.04, 462.45, 462.45, CHAVE_A)


def test_leitor_portal_filtra_competencia_cancelada_e_retencao():
    lado = co.ler_portal(io.BytesIO(xlsx_portal([
        (493, "09/2026", "9,249.04", 462.45, "2 - Retido pelo Tomador", "Normal", CHAVE_A),
        (121, "08/2026", 450.0, 0, "1 - Não Retido", "Normal", CHAVE_D),
        (5, "09/2026", 10.0, 0, "1 - Não Retido", "Cancelada", CHAVE_B),
    ])), "09/2026")
    assert lado.qtd == 1 and lado.canceladas == 1
    assert [n.numero for n in lado.outras_competencias] == [121]
    n = lado.notas[0]
    assert n.valor == pytest.approx(9249.04) and n.iss_retido == pytest.approx(462.45) and n.chave == CHAVE_A


def test_leitor_portal_recusa_planilha_que_nao_e_do_portal():
    wb = openpyxl.Workbook()
    wb.active.append(["a", "b"])
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(ValueError, match="faltam as colunas"):
        co.ler_portal(io.BytesIO(buf.getvalue()), "09/2026")


# ---------------------------------------------------------------- comparação

def lado_dec(notas):
    return co.ler_decweb(xlsx_decweb(notas))


def lado_por(notas):
    return co.ler_portal(io.BytesIO(xlsx_portal(notas)), "09/2026")


def test_bate_pela_chave_mesmo_com_numero_diferente():
    d = lado_dec([("202610000000000493", 9249.04, 462.45, 462.45, "Normal", CHAVE_A)])
    p = lado_por([(999999, "09/2026", 9249.04, 462.45, "2 - Retido pelo Tomador", "Normal", CHAVE_A)])
    assert co.comparar(d, p).situacao == "ok"


def test_sem_chave_casa_pelo_numero():
    d = lado_dec([("202610000000000493", 100.0, 5.0, 0, "Normal", "")])
    p = lado_por([(493, "09/2026", 100.0, 5.0, "1 - Não Retido", "Normal", "")])
    assert co.comparar(d, p).situacao == "ok"


def test_divergencias_nota_de_um_lado_valor_e_retencao():
    d = lado_dec([
        ("202610000000000493", 100.0, 5.0, 5.0, "Normal", CHAVE_A),
        ("202610000000000494", 200.0, 10.0, 0, "Normal", CHAVE_B),
        ("202610000000000995", 50.0, 0, 0, "Normal", CHAVE_C),
    ])
    p = lado_por([
        (493, "09/2026", 100.0, 5.0, "1 - Não Retido", "Normal", CHAVE_A),   # retenção diferente
        (494, "09/2026", 250.0, 10.0, "1 - Não Retido", "Normal", CHAVE_B),  # valor diferente
        (777, "09/2026", 9.0, 0, "1 - Não Retido", "Normal", CHAVE_D),        # só no Portal
    ])
    c = co.comparar(d, p)
    assert c.situacao == "divergencia"
    assert [n.numero for n in c.so_decweb] == [995] and [n.numero for n in c.so_portal] == [777]
    assert len(c.valor_diferente) == 1 and len(c.retido_diferente) == 1


def test_atencao_para_iss_diferente_nota_zerada_e_outra_competencia():
    d = lado_dec([("202610000000000493", 100.0, 0.0, 0, "Normal", CHAVE_A),   # Simples: DecWeb zera o ISS
                  ("202610000000000494", 0.0, 0, 0, "Normal", CHAVE_B)])
    p = lado_por([(493, "09/2026", 100.0, 2.0, "1 - Não Retido", "Normal", CHAVE_A),
                  (494, "09/2026", 0.0, 0, "1 - Não Retido", "Normal", CHAVE_B),
                  (121, "08/2026", 450.0, 0, "1 - Não Retido", "Normal", CHAVE_D)])
    c = co.comparar(d, p)
    assert c.situacao == "atencao" and len(c.iss_diferente) == 1 and len(c.zeradas) == 1


# ---------------------------------------------------------------- por cliente (com arquivos em pastas)

def test_cliente_tomados_bate_e_prestados_sem_portal(pastas):
    mun, nac = pastas
    grava_decweb(mun, "tomados", xlsx_decweb([("202610000000000493", 9249.04, 462.45, 462.45, "Normal", CHAVE_A)]))
    grava_portal(nac, "tomados", xlsx_portal([
        (493, "09/2026", 9249.04, 462.45, "2 - Retido pelo Tomador", "Normal", CHAVE_A),
        (121, "08/2026", 450.0, 0, "1 - Não Retido", "Normal", CHAVE_D)]))
    grava_decweb(mun, "prestados", xlsx_decweb([("202610000000000001", 10.0, 0.5, 0, "Normal", CHAVE_B)], "Prestados"))
    r = co.conciliar_cliente(CLIENTE, "09/2026")
    assert r["lados"]["tomados"]["situacao"] == "atencao"          # a nota de 08/2026 do Portal é só informativa
    assert "outra competência" in r["lados"]["tomados"]["texto"]
    assert r["lados"]["prestados"]["situacao"] == "sem_portal"
    assert r["lados"]["declaracao"]["situacao"] == "sem_dados"


def test_cliente_le_o_decweb_do_zip(pastas):
    mun, nac = pastas
    grava_decweb(mun, "tomados", xlsx_decweb([("202610000000000493", 100.0, 5.0, 0, "Normal", CHAVE_A)]), como_zip=True)
    grava_portal(nac, "tomados", xlsx_portal([(493, "09/2026", 100.0, 5.0, "1 - Não Retido", "Normal", CHAVE_A)]))
    assert co.conciliar_cliente(CLIENTE, "09/2026")["lados"]["tomados"]["situacao"] == "ok"


def test_portal_com_notas_e_sem_arquivo_do_decweb_diverge(pastas):
    _mun, nac = pastas
    grava_portal(nac, "prestados", xlsx_portal([(1, "09/2026", 10.0, 0, "1 - Não Retido", "Normal", CHAVE_A)]))
    assert co.conciliar_cliente(CLIENTE, "09/2026")["lados"]["prestados"]["situacao"] == "divergencia"


def test_cliente_sem_nenhum_arquivo(pastas):
    r = co.conciliar_cliente({"numero": "999", "apelido": "Ninguem"}, "09/2026")
    assert {l["situacao"] for l in r["lados"].values()} == {"sem_dados"}


def test_csv_nao_vaza_nada_alem_do_resumo(pastas, tmp_path):
    mun, nac = pastas
    grava_decweb(mun, "tomados", xlsx_decweb([("202610000000000493", 100.0, 5.0, 0, "Normal", CHAVE_A)]))
    grava_portal(nac, "tomados", xlsx_portal([(493, "09/2026", 100.0, 5.0, "1 - Não Retido", "Normal", CHAVE_A)]))
    co.gravar_csv([co.conciliar_cliente(CLIENTE, "09/2026")], str(tmp_path / "s.csv"))
    linhas = (tmp_path / "s.csv").read_text(encoding="utf-8-sig").splitlines()
    assert linhas[0].startswith("numero;apelido;lado;situacao")
    assert len(linhas) == 4 and any(";tomados;ok;1;100.0;5.0;0.0;1;100.0;5.0;0.0;" in l for l in linhas)


# ---------------------------------------------------------------- declaração e guia (texto dos PDFs reais)

TEXTO_DECLARACAO = (
    "Prefeitura de Porto Alegre Declaração Mensal - ISSQN Set/2026 Pág.1 /1 NFSE Nota Fiscal Eletrônica "
    "Resumo de Cálculo Base de Cálculo Receita Bruta 21.051,25 Deduções Legais 0,00 Base de cálculo 21.051,25 "
    "Total do imposto devido 842,02 Total Geral A Recolher 842,02 "
    "RECIBO DE ENTREGA - DECLARAÇÃO MENSAL DECLARAÇÃO RECEBIDA EM\n03/10/2026 às 11:18:03 AUTENTICAÇÃO F7 19"
)
TEXTO_DECLARACAO_DUAS_ESCRITURACOES = (
    "Base de Cálculo Receita Bruta 0,00 Deduções Legais 0,00 ... Total do imposto devido 0,00 "
    "Pág.2 Base de Cálculo Receita Bruta 48.063,00 Deduções Legais 0,00 ... Total Geral A Recolher 2.403,15 "
    "DECLARAÇÃO RECEBIDA EM 03/10/2026 às 11:24:44"
)
TEXTO_GUIA = "IMPOSTO DEVIDO R$ 842,02 NÃO RECEBER ESTA GUIA APÓS 13/10/2026ISSQN-e 03/10/2026 11:18 Até 13/10/2026: VALOR A PAGAR R$ 842,02"


def test_extrai_declaracao_guia_e_recibo():
    assert co.extrair_declaracao(TEXTO_DECLARACAO) == {"receita": 21051.25, "a_recolher": 842.02, "recibo": "03/10/2026 11:18:03"}
    assert co.extrair_declaracao(TEXTO_DECLARACAO_DUAS_ESCRITURACOES)["receita"] == pytest.approx(48063.0)
    assert co.extrair_guia(TEXTO_GUIA) == {"valor": 842.02, "vencimento": "13/10/2026"}
    assert co.extrair_declaracao("")["recibo"] is None


def test_declaracao_contra_prestados_e_guia(pastas, monkeypatch):
    mun, _nac = pastas
    base = co.nome_base(CLIENTE, "09", "2026")
    (mun / f"{base}_DeclaraçãoMensal.pdf").write_bytes(b"x")
    (mun / f"{base}_ISSQN.pdf").write_bytes(b"x")
    textos = {f"{base}_DeclaraçãoMensal.pdf": TEXTO_DECLARACAO, f"{base}_ISSQN.pdf": TEXTO_GUIA}
    monkeypatch.setattr(co, "texto_pdf", lambda caminho: textos[caminho.name])

    prest = lado_dec([("202610000000000001", 21051.25, 842.02, 0, "Normal", CHAVE_A)])
    sit, texto = co.conciliar_declaracao(CLIENTE, "09", "2026", mun, prest)
    assert sit == "ok" and "= guia, venc. 13/10/2026" in texto and "recibo 03/10/2026 11:18:03" in texto

    prest_errado = lado_dec([("202610000000000001", 20000.0, 800.0, 0, "Normal", CHAVE_A)])
    assert co.conciliar_declaracao(CLIENTE, "09", "2026", mun, prest_errado)[0] == "atencao"

    textos[f"{base}_ISSQN.pdf"] = TEXTO_GUIA.replace("VALOR A PAGAR R$ 842,02", "VALOR A PAGAR R$ 800,00")
    assert co.conciliar_declaracao(CLIENTE, "09", "2026", mun, prest)[0] == "divergencia"


def test_declaracao_sem_pypdf_nao_quebra(pastas, monkeypatch):
    mun, _nac = pastas
    (mun / f"{co.nome_base(CLIENTE, '09', '2026')}_DeclaraçãoMensal.pdf").write_bytes(b"x")
    monkeypatch.setattr(co, "texto_pdf", lambda caminho: None)
    sit, texto = co.conciliar_declaracao(CLIENTE, "09", "2026", mun, None)
    assert sit == "sem_dados" and "pypdf" in texto


def test_pdf_fora_do_formato_nao_fica_em_silencio(pastas, monkeypatch):
    mun, _nac = pastas
    (mun / f"{co.nome_base(CLIENTE, '09', '2026')}_DeclaraçãoMensal.pdf").write_bytes(b"x")
    monkeypatch.setattr(co, "texto_pdf", lambda caminho: "texto qualquer sem os campos")
    sit, texto = co.conciliar_declaracao(CLIENTE, "09", "2026", mun, None)
    assert sit == "atencao" and "NÃO consegui ler a receita bruta e o total a recolher" in texto


def test_relatorio_nota_a_nota_e_sem_pdf(pastas, tmp_path):
    mun, nac = pastas
    grava_decweb(mun, "tomados", xlsx_decweb([
        ("202610000000000493", 100.0, 5.0, 5.0, "Normal", CHAVE_A),      # ok
        ("202610000000000494", 200.0, 10.0, 0, "Normal", CHAVE_B),       # valor diferente
        ("202610000000000995", 50.0, 0, 0, "Normal", CHAVE_C),           # só no DecWeb
    ]), como_zip=True)                                                    # relação lida direto do ZIP da prefeitura
    grava_portal(nac, "tomados", xlsx_portal([
        (493, "09/2026", 100.0, 5.0, "2 - Retido pelo Tomador", "Normal", CHAVE_A),
        (494, "09/2026", 250.0, 10.0, "1 - Não Retido", "Normal", CHAVE_B),
        (777, "09/2026", 9.0, 0, "1 - Não Retido", "Normal", CHAVE_D),
        (121, "08/2026", 450.0, 0, "1 - Não Retido", "Normal", "9" * 50),
    ]))
    r = co.conciliar_cliente(CLIENTE, "09/2026", com_pdf=False)
    assert r["lados"]["declaracao"]["texto"] == "PDF não conferido (--sem-pdf)"
    co.gravar_csv_notas([r], str(tmp_path / "n.csv"))
    linhas = [l.split(";") for l in (tmp_path / "n.csv").read_text(encoding="utf-8-sig").splitlines()[1:]]
    por_nota = {l[4]: l[3] for l in linhas}
    assert por_nota == {"493": "ok", "494": "valor_diferente", "995": "so_no_decweb", "777": "so_no_portal",
                        "121": "portal_outra_competencia_08-2026"}


def test_nota_de_substituicao_gerada_conta_como_normal():
    from conciliacao import Nota
    assert not Nota(1, 10.0, situacao="101 - NFS-e de Substituição Gerada Normal").cancelada
    assert Nota(2, 10.0, situacao="100 - NFS-e Gerada Substituída").cancelada
    assert Nota(3, 10.0, situacao="100 - NFS-e Gerada Cancelada").cancelada
